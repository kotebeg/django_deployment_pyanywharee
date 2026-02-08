from io import BytesIO

from django.test import TestCase
from django.urls import reverse, resolve
from django.contrib.auth.models import User
from rest_framework.test import APIClient
from rest_framework import status

import openpyxl


class TokenAuthTests(TestCase):
    """Tests for POST /api/v1/auth/token/"""

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username='apiuser', password='apipass123'
        )
        self.url = reverse('api:api_token')

    def test_obtain_token_valid_credentials(self):
        response = self.client.post(self.url, {
            'username': 'apiuser',
            'password': 'apipass123',
        })
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('token', response.data)

    def test_obtain_token_invalid_credentials(self):
        response = self.client.post(self.url, {
            'username': 'apiuser',
            'password': 'wrongpass',
        })
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_obtain_token_missing_fields(self):
        response = self.client.post(self.url, {})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)


class GenerateSignalAPITests(TestCase):
    """Tests for GET /api/v1/signals/generate/"""

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username='apiuser', password='apipass123'
        )
        self.url = reverse('api:api_signal_generate')

    def test_generate_requires_auth(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_generate_returns_signal_data(self):
        self.client.force_authenticate(user=self.user)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('num_points', response.data)
        self.assertIn('data', response.data)
        self.assertEqual(response.data['num_points'], 700)
        self.assertEqual(len(response.data['data']), 700)

    def test_generate_data_structure(self):
        self.client.force_authenticate(user=self.user)
        response = self.client.get(self.url)
        first_point = response.data['data'][0]
        self.assertIn('time', first_point)
        self.assertIn('random_signal', first_point)

    def test_generate_with_token_header(self):
        token_response = self.client.post(reverse('api:api_token'), {
            'username': 'apiuser',
            'password': 'apipass123',
        })
        token = token_response.data['token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {token}')
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_generate_xlsx_format(self):
        self.client.force_authenticate(user=self.user)
        response = self.client.get(self.url, {'output': 'xlsx'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment', response['Content-Disposition'])
        wb = openpyxl.load_workbook(BytesIO(response.content))
        self.assertIn('RandomSignal', wb.sheetnames)

    def test_generate_url_resolves(self):
        view = resolve('/api/v1/signals/generate/')
        self.assertEqual(view.url_name, 'api_signal_generate')


class UploadSignalAPITests(TestCase):
    """Tests for POST /api/v1/signals/upload/"""

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username='apiuser', password='apipass123'
        )
        self.url = reverse('api:api_signal_upload')

    def _make_excel_file(self, num_rows=10):
        """Create a valid Excel file in memory for testing."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'RandomSignal'
        for i in range(num_rows):
            ws.cell(row=i + 1, column=1, value=i)
            ws.cell(row=i + 1, column=2, value=i * 0.5)
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = 'test_signal.xlsx'
        return buf

    def test_upload_requires_auth(self):
        response = self.client.post(self.url)
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_upload_valid_excel(self):
        self.client.force_authenticate(user=self.user)
        excel_file = self._make_excel_file(10)
        response = self.client.post(
            self.url,
            {'file': excel_file},
            format='multipart'
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['num_points'], 10)
        self.assertEqual(len(response.data['data']), 10)

    def test_upload_no_file(self):
        self.client.force_authenticate(user=self.user)
        response = self.client.post(self.url)
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('error', response.data)

    def test_upload_data_values(self):
        self.client.force_authenticate(user=self.user)
        excel_file = self._make_excel_file(5)
        response = self.client.post(
            self.url,
            {'file': excel_file},
            format='multipart'
        )
        first = response.data['data'][0]
        self.assertEqual(first['time'], 0)
        self.assertAlmostEqual(first['random_signal'], 0.0)

    def test_upload_url_resolves(self):
        view = resolve('/api/v1/signals/upload/')
        self.assertEqual(view.url_name, 'api_signal_upload')
