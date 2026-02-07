from io import BytesIO

from django.test import TestCase, Client
from django.urls import reverse, resolve
from django.contrib.auth.models import User

import pandas as pd
import openpyxl

from signal_viz import views
from signal_viz.signal_generator import (
    generate_random_sequence,
    generate_random_signal_plot,
)


class SignalGeneratorTests(TestCase):
    def test_generate_random_sequence_returns_dataframe(self):
        df = generate_random_sequence()
        self.assertIsInstance(df, pd.DataFrame)

    def test_generate_random_sequence_columns(self):
        df = generate_random_sequence()
        self.assertListEqual(list(df.columns), ["time", "Random Signal"])

    def test_generate_random_sequence_length(self):
        df = generate_random_sequence()
        self.assertEqual(len(df), 700)

    def test_generate_random_signal_plot_returns_html(self):
        df = generate_random_sequence()
        html = generate_random_signal_plot(df)
        self.assertIsInstance(html, str)
        self.assertIn("plotly", html.lower())


class SignalVizViewTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.upload_url = reverse("signal_viz_home")
        self.user = User.objects.create_user(
            username="testuser", password="testpass123"
        )

    def test_signal_viz_requires_auth(self):
        response = self.client.get(self.upload_url)
        self.assertRedirects(response, reverse("login"))

    def test_signal_viz_get_renders_plot(self):
        self.client.login(username="testuser", password="testpass123")
        response = self.client.get(self.upload_url)
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "signal_viz/index.html")
        self.assertIn("plot", response.context)

    def test_signal_viz_post_requires_auth(self):
        response = self.client.post(self.upload_url)
        self.assertRedirects(response, reverse("login"))

    def test_signal_viz_post_valid_excel(self):
        self.client.login(username="testuser", password="testpass123")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "RandomSignal"
        for i in range(10):
            ws.cell(row=i + 1, column=1, value=i)
            ws.cell(row=i + 1, column=2, value=i * 0.5)
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = "test_signal.xlsx"

        response = self.client.post(
            self.upload_url,
            {"uploaded_file": buf},
            format="multipart",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("plot", response.context)

    def test_signal_viz_url_resolves(self):
        view = resolve("/signal_viz/")
        self.assertEqual(view.func.view_class, views.ProfileView)


class ExcelDownloadTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            username="testuser", password="testpass123"
        )

    def test_get_excel_returns_xlsx(self):
        self.client.login(username="testuser", password="testpass123")
        # Hit the upload page first to populate session with df_data
        self.client.get(reverse("signal_viz_home"))

        response = self.client.get(reverse("get_excel"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/ms-excel")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])

    def test_get_excel_contains_data(self):
        self.client.login(username="testuser", password="testpass123")
        self.client.get(reverse("signal_viz_home"))

        response = self.client.get(reverse("get_excel"))
        wb = openpyxl.load_workbook(BytesIO(response.content))
        self.assertIn("RandomSignal", wb.sheetnames)
        ws = wb["RandomSignal"]
        self.assertGreater(ws.max_row, 0)

    def test_get_excel_url_resolves(self):
        view = resolve("/signal_viz/get_excel")
        self.assertEqual(view.func, views.get_excel)
