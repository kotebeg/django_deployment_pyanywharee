from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from rest_framework.parsers import MultiPartParser

from django.http import HttpResponse

import pandas as pd
import openpyxl
from io import BytesIO
from datetime import datetime

from .signal_generator import generate_random_sequence
from .serializers import SignalResponseSerializer


class GenerateSignalView(APIView):
    """
    GET /api/v1/signals/generate/
    Generate a random sinusoidal signal with Gaussian noise.

    Query parameters:
        output: 'json' (default) or 'xlsx'
    """

    def get(self, request):
        df = generate_random_sequence()

        if request.query_params.get('output') == 'xlsx':
            return self._excel_response(df)

        data_dicts = df.to_dict('records')
        serializer = SignalResponseSerializer({
            'num_points': len(data_dicts),
            'data': data_dicts,
        })
        return Response(serializer.data)

    def _excel_response(self, df):
        now = datetime.now().strftime("%Y%m%d_%H%M%S")
        workbook = openpyxl.Workbook()
        sheet = workbook.create_sheet(title='RandomSignal')
        del workbook['Sheet']

        for _, row in df.iterrows():
            x_val = int(row['time'])
            y_val = row['Random Signal']
            sheet.cell(row=x_val + 1, column=1, value=x_val)
            sheet.cell(row=x_val + 1, column=2, value=y_val)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="signal_{now}.xlsx"'
        return response


class UploadSignalView(APIView):
    """
    POST /api/v1/signals/upload/
    Upload an Excel file containing signal data and return parsed JSON.

    Expects a file field named 'file' with a sheet named 'RandomSignal'.
    """
    parser_classes = [MultiPartParser]

    def post(self, request):
        if 'file' not in request.FILES:
            return Response(
                {'error': 'No file provided. Send a file in the "file" field.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        uploaded_file = request.FILES['file']

        try:
            df = pd.read_excel(
                uploaded_file,
                sheet_name='RandomSignal',
                header=None,
                names=['time', 'Random Signal']
            )
        except Exception as e:
            return Response(
                {'error': f'Failed to parse Excel file: {str(e)}'},
                status=status.HTTP_400_BAD_REQUEST
            )

        data_dicts = df.to_dict('records')
        serializer = SignalResponseSerializer({
            'num_points': len(data_dicts),
            'data': data_dicts,
        })
        return Response(serializer.data)
