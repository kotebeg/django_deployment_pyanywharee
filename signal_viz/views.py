from django.shortcuts import render
from django.http import HttpResponse, HttpResponseRedirect
from django.urls import reverse
from django.views import View

from .signal_generator import generate_random_signal_plot, generate_random_sequence
from .forms import excel_upload_form

import pandas as pd
import openpyxl
from io import StringIO
from datetime import datetime


class ProfileView(View):
    def get(self, request):
        if not request.user.is_authenticated:
            return HttpResponseRedirect(reverse("login"))

        excel_upload = excel_upload_form()
        self.df_data = generate_random_sequence()
        request.session['df_data'] = generate_random_sequence().to_json()
        plot_for_front = generate_random_signal_plot(self.df_data)

        return render(request, 'signal_viz/index.html', {
            'var1': 'var_from_front',
            'form': excel_upload,
            'plot': plot_for_front,
        })

    def post(self, request):
        if not request.user.is_authenticated:
            return HttpResponseRedirect(reverse("login"))

        excel_upload = excel_upload_form()
        uploaded_file_from_front = request.FILES['uploaded_file']
        uploaded_file_from_front_df = pd.read_excel(
            uploaded_file_from_front,
            sheet_name='RandomSignal',
            header=None,
            names=["time", "Random Signal"]
        )
        request.session['df_data'] = uploaded_file_from_front_df.to_json()
        plot_for_front_from_imported_data = generate_random_signal_plot(uploaded_file_from_front_df)

        return render(request, 'signal_viz/index.html', {
            'var1': 'var_from_front',
            'form': excel_upload,
            'plot': plot_for_front_from_imported_data,
        })


def get_excel(request):
    now = datetime.now().strftime("%Y %m %d %H %M %S")
    workbook = openpyxl.Workbook()

    response = HttpResponse(content_type='application/ms-excel')
    response['Content-Disposition'] = f'attachment; filename="random_signal_{now}.xlsx"'

    sector_data_sheet = workbook.create_sheet(title='RandomSignal')
    del workbook["Sheet"]

    df = pd.read_json(StringIO(request.session['df_data']))

    for i in df.iloc:
        x_val = int(i.values[0])
        y_val = i.values[1]
        sector_data_sheet.cell(row=x_val+1, column=1, value=x_val)
        sector_data_sheet.cell(row=x_val+1, column=2, value=y_val)

    workbook.save(response)

    return response
