from django.shortcuts import render

from . signal_generator import generate_random_signal_plot, generate_random_sequence

from django.contrib.auth import authenticate, login, logout
import pandas as pd

import openpyxl
from openpyxl.styles import Font
from io import StringIO

# activity
from datetime import datetime
now = datetime.now().strftime("%Y %m %d %H %M %S")

from django.http import HttpResponseRedirect
from django.urls import reverse

# Create your views here.
from django.http import HttpResponse, HttpResponseNotFound, HttpResponseRedirect

from .forms import excel_upload_form

from django.views import View

# plot_for_front = None
# df_data= None
class ProfileView(View):
        # def __init__(self,):
                # self.df_data =  generate_random_sequence()
        def get(self, request):
                if not request.user.is_authenticated:
                        return HttpResponseRedirect(reverse("login"))
                else:
                        excel_upload = excel_upload_form()
                        # global plot_for_front
                        # global df_data
                        # self.df_data = generate_random_sequence()
                        request.session['df_data'] = generate_random_sequence().to_json()
                        plot_for_front = generate_random_signal_plot(pd.read_json(request.session['df_data']))

                        # print('--->Globals',globals())

                        return render(request, 'file_upload/index.html', {
                        'var1': 'var_from_front',
                        'form': excel_upload,
                        'plot': plot_for_front,
                        })

        def post(self, request):
                if not request.user.is_authenticated:
                        return HttpResponseRedirect(reverse("login"))
                else:

                        excel_upload = excel_upload_form()
                        uploaded_excel = excel_upload_form(request.POST, request.FILES)

                        uploaded_file_from_front = request.FILES['uploaded_file']
                        uploaded_file_from_front_df = pd.read_excel(uploaded_file_from_front, sheet_name='RandomSignal', header=None, names=["time","Random Signal"])
                        # uploaded_file_from_front_df = uploaded_file_from_front_df.set_index(uploaded_file_from_front_df.columns[0])
                        # print('---->xxxhere')
                        # print(uploaded_file_from_front_df)
                        # print(uploaded_file_from_front_df)
                        request.session['df_data'] = uploaded_file_from_front_df.to_json()
                        # print(uploaded_file_from_front_df)

                        plot_for_front_from_imported_data = generate_random_signal_plot(uploaded_file_from_front_df)
                        # print('file uploaded --->', request.FILES['uploaded_file'])

                        # if uploaded_excel.is_valid():
                        #         request.session['file_uploaded'] = True

                        return render(request, 'file_upload/index.html', {
                        'var1': 'var_from_front',
                        'form': excel_upload,
                        'plot': plot_for_front_from_imported_data,
                        
                        })

def get_excel(request):
        # Create a new workbook
        # print('--->> download excell', type(df_data))

        # print(plot_for_front)
        # for row in df_data.iloc:
        #         print(row)
        # print('--->Globals',globals())
        workbook = openpyxl.Workbook()

        response = HttpResponse(content_type='application/ms-excel')
        response['Content-Disposition'] = f'attachment; filename="random_signal_{now}.xlsx"'


        sector_data_sheet = workbook.create_sheet(title='RandomSignal')
        del workbook["Sheet"]


        # print('----++>>', pd.read_json(request.session['df_data']))



        df = pd.read_json(StringIO(request.session['df_data']))

        for i in df.iloc:
                # i.values
                x_val = int(i.values[0])
                y_val = i.values[1]
                # print(x_val, y_val)
                cell = sector_data_sheet.cell(row=x_val+1, column = 1, value = x_val)
                cell = sector_data_sheet.cell(row=x_val+1, column = 2, value = y_val)

        workbook.save(response)

        return response  